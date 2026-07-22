"""
Build Trustworthy AI KSS deck v2.
Architecture narrative: Capability → Failure → Control → Evidence → Governance.
Narrative anchor: DocuBot — an enterprise AI agent processing confidential documents.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR_TYPE
import math

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# ── Colors ───────────────────────────────────────────────────────────
NAVY       = RGBColor(0x1B, 0x2A, 0x4A)
TEAL       = RGBColor(0x00, 0x89, 0x7B)
AMBER      = RGBColor(0xFF, 0x8F, 0x00)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT   = RGBColor(0x2D, 0x37, 0x48)
LIGHT_BG   = RGBColor(0xF5, 0xF7, 0xFA)
CORAL      = RGBColor(0xE5, 0x3E, 0x3E)
MID_GRAY   = RGBColor(0xA0, 0xAE, 0xC0)
TEAL_LIGHT = RGBColor(0xE0, 0xF2, 0xF1)
NAVY_LIGHT = RGBColor(0xE8, 0xEA, 0xF0)
PURPLE     = RGBColor(0x6B, 0x46, 0xC1)
GREEN      = RGBColor(0x38, 0xA1, 0x69)
BLUE_STEEL = RGBColor(0x4A, 0x55, 0x68)
AMBER_LIGHT= RGBColor(0xFF, 0xF8, 0xE1)
RED_BG     = RGBColor(0xFD, 0xE8, 0xE8)
GREEN_BG   = RGBColor(0xE8, 0xF5, 0xE9)
BLUE_BG    = RGBColor(0xE3, 0xF2, 0xFD)

# ── Helpers ──────────────────────────────────────────────────────────
def S(): return prs.slides.add_slide(prs.slide_layouts[6])

def N(slide, text):
    slide.notes_slide.notes_text_frame.text = text

def TB(slide, left, top, width, height, text="", fs=18, bold=False, color=DARK_TEXT,
       align=PP_ALIGN.LEFT, ls=1.15):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    txBox.text_frame.word_wrap = True
    p = txBox.text_frame.paragraphs[0]
    p.text = text; p.font.size = Pt(fs); p.font.bold = bold
    p.font.color.rgb = color; p.font.name = "Calibri"; p.alignment = align
    p.line_spacing = Pt(fs * ls); p.space_after = Pt(4)
    return txBox

def ML(slide, left, top, width, height, lines, fs=14, color=DARK_TEXT, ls=1.35):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    txBox.text_frame.word_wrap = True
    for i, line in enumerate(lines):
        p = txBox.text_frame.paragraphs[0] if i == 0 else txBox.text_frame.add_paragraph()
        p.text = line; p.font.size = Pt(fs); p.font.color.rgb = color
        p.font.name = "Calibri"; p.line_spacing = Pt(fs * ls); p.space_after = Pt(3)
    return txBox

def R(slide, left, top, width, height, fill=NAVY, border=None, radius=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid(); shape.fill.fore_color.rgb = fill
    if border:
        shape.line.color.rgb = border; shape.line.width = Pt(1)
    else:
        shape.line.fill.background()
    return shape

def TITLE(slide, title, subtitle=None, color=NAVY):
    R(slide, 0, 0, 13.333, 1.35, fill=color)
    TB(slide, 0.8, 0.2, 11.7, 0.7, title, fs=30, bold=True, color=WHITE)
    if subtitle:
        TB(slide, 0.8, 0.75, 11.7, 0.45, subtitle, fs=14, color=RGBColor(0xCC, 0xD5, 0xE0))
    R(slide, 0, 1.35, 13.333, 0.04, fill=TEAL)

def PN(slide, num):
    TB(slide, 11.8, 7.05, 1.2, 0.35, str(num), fs=10, color=MID_GRAY, align=PP_ALIGN.RIGHT)

def DIV(slide, part, title, subtitle=""):
    R(slide, 0, 0, 13.333, 7.5, fill=NAVY)
    TB(slide, 0.8, 2.0, 11.7, 0.5, f"PART {part}", fs=16, color=TEAL, bold=True)
    TB(slide, 0.8, 2.5, 11.7, 1.2, title, fs=40, bold=True, color=WHITE)
    if subtitle:
        TB(slide, 0.8, 3.7, 11.7, 0.8, subtitle, fs=18, color=RGBColor(0xCC, 0xD5, 0xE0))
    R(slide, 0.8, 5.0, 2.0, 0.05, fill=TEAL)

def BADGE(slide, left, top, width, height, label, color):
    shape = R(slide, left, top, width, height, fill=color, radius=0.06)
    tf = shape.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = label; p.font.size = Pt(10)
    p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER

def BOX(slide, left, top, w, h, title, lines, title_color=TEAL, fs=12, tfs=14):
    R(slide, left, top, w, h, fill=WHITE, border=RGBColor(0xE2, 0xE8, 0xF0), radius=0.08)
    R(slide, left, top, w, 0.45, fill=title_color, radius=0.06)
    TB(slide, left+0.15, top+0.05, w-0.3, 0.35, title, fs=tfs, bold=True, color=WHITE)
    ML(slide, left+0.15, top+0.6, w-0.3, h-0.75, lines, fs=fs, color=DARK_TEXT, ls=1.35)

def ICON(slide, x, y, size, label, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(size), Inches(size))
    shape.fill.solid(); shape.fill.fore_color.rgb = color; shape.line.fill.background()
    tf = shape.text_frame; tf.word_wrap = False
    p = tf.paragraphs[0]; p.text = label; p.font.size = Pt(16)
    p.font.bold = True; p.font.color.rgb = WHITE; p.font.name = "Calibri"
    p.alignment = PP_ALIGN.CENTER

